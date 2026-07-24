import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")

for i in range(6, 55):
    slide = prs.slides[i]
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame and shape.text.strip() == "온라인 최저가":
            val_shape = slide.shapes[idx+2]
            if val_shape.shape_type != 17:  # 17 is TEXT_BOX
                print(f"Slide {i+1} WARNING: idx+2 is type {val_shape.shape_type}, not TEXT_BOX!")
                print(f"  Shape idx+1: type {slide.shapes[idx+1].shape_type}, text: {repr(slide.shapes[idx+1].text)}")
                print(f"  Shape idx+2: type {slide.shapes[idx+2].shape_type}, text: {repr(slide.shapes[idx+2].text)}")
                if idx+3 < len(slide.shapes):
                    print(f"  Shape idx+3: type {slide.shapes[idx+3].shape_type}, text: {repr(slide.shapes[idx+3].text)}")
