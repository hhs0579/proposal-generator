import json
import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")

products = []

for i, slide in enumerate(prs.slides):
    if i < 6 or i > 54:
        continue
    
    product_name = None
    options_text = ""
    
    # Extract product name
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame and "KEY SELLING POINTS" in shape.text:
            if idx > 0 and slide.shapes[idx-1].has_text_frame:
                product_name = slide.shapes[idx-1].text.strip().replace('\n', ' ')
            break
            
    if not product_name:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip().replace('\n', ' ')
                if text and len(text) > 2 and len(text) < 50 and "상품 제안서" not in text and "PRODUCT LINE-UP" not in text:
                    product_name = text
                    break
    
    if not product_name:
        product_name = f"미지정 상품 (슬라이드 {i+1})"

    # Extract options text
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame and shape.text.strip() == "옵션 및 구성":
            if idx + 2 < len(slide.shapes) and slide.shapes[idx+2].has_text_frame:
                options_text = slide.shapes[idx+2].text.strip()
            break
            
    products.append({
        "slide_index": i,
        "product_name": product_name,
        "options": options_text
    })

with open("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/products_rich.json", "w", encoding="utf-8") as f:
    json.dump(products, f, ensure_ascii=False, indent=2)

print("Extracted", len(products), "products")
