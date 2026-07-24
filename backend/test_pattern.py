import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")

for i in range(6, 11):
    slide = prs.slides[i]
    print(f"Slide {i}")
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = shape.text.strip()
            if text in ["온라인 최저가", "공급가", "재고수량", "배송비", "카톤수량", "옵션 및 구성"]:
                print(f"  {text} is at {idx}, value box should be {idx+2}")
                if idx+2 < len(slide.shapes) and slide.shapes[idx+2].has_text_frame:
                    print(f"    Value box exists: {repr(slide.shapes[idx+2].text)}")
                else:
                    print("    Value box NOT FOUND or NO TEXT FRAME")
