import sys
from pptx import Presentation
from pptx.util import Pt

def replace_text_in_shape(shape, new_text):
    if not shape.has_text_frame: return
    text_frame = shape.text_frame
    
    if not text_frame.paragraphs:
        return
        
    p0 = text_frame.paragraphs[0]
    if not p0.runs:
        p0.add_run()
        
    p0.runs[0].text = str(new_text)
    p0.runs[0].font.size = Pt(9)
    
    for r in p0.runs[1:]:
        r.text = ""
        
    for p in text_frame.paragraphs[1:]:
        p_element = p._p
        p_element.getparent().remove(p_element)

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")
slide = prs.slides[7]
shape = slide.shapes[46] # 공급가 value box

print("BEFORE:")
print(f"Alignment: {shape.text_frame.paragraphs[0].alignment}")

replace_text_in_shape(shape, "테스트 가격\n줄바꿈 테스트")

print("\nAFTER:")
print(f"Alignment: {shape.text_frame.paragraphs[0].alignment}")

prs.save("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/test_overlap.pptx")
