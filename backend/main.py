import io
import os
import json
import base64
import uuid
import traceback
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles

from pydantic import BaseModel
from typing import List, Optional
from pptx import Presentation
from pptx.util import Pt

app = FastAPI()
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
PRODUCTS_PATH = os.path.join(BASE_DIR, "products_rich.json")
CUSTOM_PRODUCTS_PATH = os.path.join(DATA_DIR, "custom_products.json")
PRODUCT_OVERRIDES_PATH = os.path.join(DATA_DIR, "product_overrides.json")
HIDDEN_PRODUCTS_PATH = os.path.join(DATA_DIR, "hidden_products.json")
DRAFT_PATH = os.path.join(DATA_DIR, "saved_draft.json")
TEMPLATE_PATH = os.path.join(BASE_DIR, "template.pptx")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

class SelectedProduct(BaseModel):
    id: Optional[int] = None
    catalog_id: Optional[str] = None
    slide_index: int
    name: Optional[str] = ""
    description: str = ""
    detailLink: str = ""
    onlineLowestPrice: str = ""
    supplyPrice: str = ""
    stockQuantity: str = ""
    shippingFee: str = ""
    cartonQuantity: str = ""
    options: str = ""
    image_base64: Optional[str] = None
    is_new: Optional[bool] = False

class ProposalRequest(BaseModel):
    products: List[SelectedProduct]

class DeleteProductRequest(BaseModel):
    catalog_id: str

def delete_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    try:
        xml_slides.remove(slides[index])
    except Exception as e:
        print(f"Failed to delete slide {index}: {e}")


def rebuild_slides_in_order(prs, ordered_indices):
    """선택한 순서대로 슬라이드를 재배치한다. (표지 → 상품들 → 마지막 장)"""
    sld_id_lst = prs.slides._sldIdLst
    slides = list(sld_id_lst)
    ordered = [slides[i] for i in ordered_indices]
    for element in list(sld_id_lst):
        sld_id_lst.remove(element)
    for element in ordered:
        sld_id_lst.append(element)

def replace_text_in_shape(shape, new_text, font_size=9, is_bold=False):
    if not shape.has_text_frame:
        return
    text_frame = shape.text_frame
    text_frame.clear()

    lines = str(new_text).replace("\x0b", "\n").split("\n")
    if not lines:
        lines = [""]

    for i, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        if is_bold:
            run.font.bold = True


B2B_LABELS = {
    "온라인 최저가",
    "공급가",
    "재고수량",
    "배송비",
    "카톤수량",
    "옵션 및 구성",
}

PROTECTED_TEXTS = B2B_LABELS | {
    "B2B 거래 검토 조건",
    "미확정 항목은 공란 처리",
    "상세페이지 링크",
    "KEY SELLING POINTS",
}


def find_value_shapes_for_label(slide, label_shape, label_idx=None):
    """라벨 오른쪽 값 칸에 해당하는 모든 텍스트 도형(중복 박스 포함)을 찾는다."""
    candidates = []
    seen = set()

    def add(shape):
        shape_id = id(shape)
        if shape_id in seen:
            return
        seen.add(shape_id)
        candidates.append(shape)

    if label_idx is not None and label_idx + 2 < len(slide.shapes):
        primary = slide.shapes[label_idx + 2]
        if getattr(primary, "has_text_frame", False):
            add(primary)

    for shape in slide.shapes:
        if shape == label_shape or not getattr(shape, "has_text_frame", False):
            continue
        try:
            text = shape.text.strip()
            if text in PROTECTED_TEXTS or text.startswith("http"):
                continue
            # 라벨 바로 오른쪽 값 영역만 (다른 열 라벨/값 제외)
            if shape.left <= label_shape.left:
                continue
            if shape.left > label_shape.left + 2800000:
                continue
            # 같은 행 근처 (옵션 박스는 약간 위아래로 여유)
            if abs(shape.top - label_shape.top) > 380000:
                continue
            add(shape)
        except Exception:
            continue

    return candidates


def set_labeled_value(slide, label_text, new_value, font_size=9):
    """
    템플릿에 있던 기존 값은 전부 지우고,
    사용자가 작성한 내용만 넣는다. 공란이면 공란 유지.
    """
    label_idx = None
    label_shape = None
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame and shape.text.strip() == label_text:
            label_idx = idx
            label_shape = shape
            break
    if label_shape is None:
        return

    targets = find_value_shapes_for_label(slide, label_shape, label_idx)
    if not targets:
        return

    # 1) 겹친/잔여 박스 전부 공란 처리
    for shape in targets:
        replace_text_in_shape(shape, "", font_size=font_size)

    # 2) 작성한 내용이 있을 때만 대표 박스에 기록
    value = "" if new_value is None else str(new_value)
    if not value.strip():
        return

    primary = targets[0]
    if label_idx is not None and label_idx + 2 < len(slide.shapes):
        candidate = slide.shapes[label_idx + 2]
        if candidate in targets:
            primary = candidate

    replace_text_in_shape(primary, value, font_size=font_size)


def set_detail_link(slide, link_value, font_size=9):
    """상세페이지 링크 라벨 다음 박스에 URL을 넣는다. 공란이면 비운다."""
    for idx, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        if shape.text.strip() != "상세페이지 링크":
            continue
        for j in range(idx + 1, min(idx + 4, len(slide.shapes))):
            target = slide.shapes[j]
            if not target.has_text_frame:
                continue
            replace_text_in_shape(target, (link_value or "").strip(), font_size=font_size)
            return


def apply_product_b2b_fields(slide, prod):
    """상품 슬라이드의 B2B 항목과 상세페이지 링크를 사용자 입력 기준으로 덮어쓴다."""
    set_labeled_value(slide, "온라인 최저가", prod.onlineLowestPrice or "")
    set_labeled_value(slide, "공급가", prod.supplyPrice or "")
    set_labeled_value(slide, "재고수량", prod.stockQuantity or "")
    set_labeled_value(slide, "배송비", prod.shippingFee or "")
    set_labeled_value(slide, "카톤수량", prod.cartonQuantity or "")
    set_labeled_value(slide, "옵션 및 구성", prod.options or "")
    set_detail_link(slide, getattr(prod, "detailLink", "") or "")


def needs_slide_assignment(product: SelectedProduct) -> bool:
    """신규/커스텀 상품은 템플릿 전용 슬라이드가 없으므로 빈 상품 슬라이드를 할당한다."""
    if product.is_new:
        return True
    if product.slide_index is None or product.slide_index < 0:
        return True
    if str(product.catalog_id or "").startswith("custom:"):
        return True
    return False


def read_json(path, default):
    if not os.path.exists(path) or os.path.isdir(path):
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def ensure_writable_json(path, default):
    directory = os.path.dirname(path)
    if directory:
        os.makedirs(directory, exist_ok=True)

    if os.path.isdir(path):
        raise OSError(f"Expected a JSON file but found a directory: {path}")

    if not os.path.exists(path):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(default, f, ensure_ascii=False, indent=2)

    try:
        os.chmod(path, 0o666)
    except OSError:
        pass


def write_json(path, data):
    default = [] if isinstance(data, list) else {}
    ensure_writable_json(path, default)
    # Docker 단일 파일 마운트에서는 os.replace가 실패할 수 있어 직접 덮어쓴다.
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    try:
        os.chmod(path, 0o666)
    except OSError:
        pass


os.makedirs(DATA_DIR, exist_ok=True)
ensure_writable_json(CUSTOM_PRODUCTS_PATH, [])
ensure_writable_json(PRODUCT_OVERRIDES_PATH, {})
ensure_writable_json(HIDDEN_PRODUCTS_PATH, [])
ensure_writable_json(DRAFT_PATH, [])


def catalog_id_for(product: SelectedProduct) -> str:
    if product.catalog_id:
        return product.catalog_id
    if product.is_new:
        return f"custom:{uuid.uuid4().hex}"
    return f"template:{product.slide_index}"


def product_snapshot(product: SelectedProduct, catalog_id: str, is_custom: bool):
    return {
        "catalog_id": catalog_id,
        "slide_index": -1 if is_custom else product.slide_index,
        "product_name": (product.name or "").strip(),
        "description": product.description,
        "detailLink": product.detailLink,
        "options": product.options,
        "onlineLowestPrice": product.onlineLowestPrice,
        "supplyPrice": product.supplyPrice,
        "stockQuantity": product.stockQuantity,
        "shippingFee": product.shippingFee,
        "cartonQuantity": product.cartonQuantity,
        "image_base64": product.image_base64,
        "is_custom": is_custom,
    }


def apply_overrides(product: dict, overrides: dict) -> dict:
    catalog_id = product.get("catalog_id")
    if not catalog_id or catalog_id not in overrides:
        return product
    merged = {**product}
    for key, value in overrides[catalog_id].items():
        if key in ("catalog_id", "slide_index", "is_custom"):
            continue
        if value is not None:
            merged[key] = value
    return merged


@app.get("/products")
async def get_products():
    try:
        overrides = read_json(PRODUCT_OVERRIDES_PATH, {})
        if not isinstance(overrides, dict):
            overrides = {}
        hidden = set(read_json(HIDDEN_PRODUCTS_PATH, []) or [])

        products = read_json(PRODUCTS_PATH, [])
        built_in_products = [
            apply_overrides(
                {
                    **product,
                    "catalog_id": f"template:{product['slide_index']}",
                    "is_custom": False,
                },
                overrides,
            )
            for product in products
            if f"template:{product['slide_index']}" not in hidden
        ]

        custom_products = [
            apply_overrides(product, overrides)
            for product in read_json(CUSTOM_PRODUCTS_PATH, [])
            if product.get("catalog_id") not in hidden
        ]
        return built_in_products + custom_products
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/products/delete")
async def delete_catalog_product(req: DeleteProductRequest):
    """기존 상품 목록에서 상품을 삭제(숨김)한다."""
    try:
        catalog_id = (req.catalog_id or "").strip()
        if not catalog_id:
            raise HTTPException(status_code=400, detail="catalog_id가 필요합니다.")

        if catalog_id.startswith("custom:"):
            custom_products = [
                item
                for item in read_json(CUSTOM_PRODUCTS_PATH, [])
                if item.get("catalog_id") != catalog_id
            ]
            write_json(CUSTOM_PRODUCTS_PATH, custom_products)
        else:
            hidden = read_json(HIDDEN_PRODUCTS_PATH, [])
            if not isinstance(hidden, list):
                hidden = []
            if catalog_id not in hidden:
                hidden.append(catalog_id)
                write_json(HIDDEN_PRODUCTS_PATH, hidden)

        overrides = read_json(PRODUCT_OVERRIDES_PATH, {})
        if isinstance(overrides, dict) and catalog_id in overrides:
            del overrides[catalog_id]
            write_json(PRODUCT_OVERRIDES_PATH, overrides)

        return {"message": "Deleted", "catalog_id": catalog_id}
    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/products/custom")
async def save_custom_product(product: SelectedProduct):
    try:
        if not product.name or not product.name.strip():
            raise HTTPException(status_code=400, detail="상품명을 입력해 주세요.")

        custom_products = read_json(CUSTOM_PRODUCTS_PATH, [])
        catalog_id = catalog_id_for(product)
        saved_product = product_snapshot(product, catalog_id, True)

        existing_index = next(
            (
                index
                for index, item in enumerate(custom_products)
                if item.get("catalog_id") == catalog_id
            ),
            None,
        )
        if existing_index is None:
            custom_products.append(saved_product)
        else:
            custom_products[existing_index] = saved_product

        write_json(CUSTOM_PRODUCTS_PATH, custom_products)
        return saved_product
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/products/persist")
async def persist_products(req: ProposalRequest):
    """상품 폼 값을 영구 저장한다. 다음 선택 시에도 가격/수량 등이 채워진다."""
    try:
        if not req.products:
            raise HTTPException(status_code=400, detail="저장할 상품이 없습니다.")

        overrides = read_json(PRODUCT_OVERRIDES_PATH, {})
        if not isinstance(overrides, dict):
            overrides = {}
        custom_products = read_json(CUSTOM_PRODUCTS_PATH, [])
        results = []

        for product in req.products:
            is_custom = bool(product.is_new) or (
                product.catalog_id or ""
            ).startswith("custom:")
            catalog_id = catalog_id_for(product)
            snapshot = product_snapshot(product, catalog_id, is_custom)

            overrides[catalog_id] = {
                "product_name": snapshot["product_name"],
                "description": snapshot["description"],
                "detailLink": snapshot["detailLink"],
                "options": snapshot["options"],
                "onlineLowestPrice": snapshot["onlineLowestPrice"],
                "supplyPrice": snapshot["supplyPrice"],
                "stockQuantity": snapshot["stockQuantity"],
                "shippingFee": snapshot["shippingFee"],
                "cartonQuantity": snapshot["cartonQuantity"],
                "image_base64": snapshot["image_base64"],
            }

            if is_custom:
                existing_index = next(
                    (
                        index
                        for index, item in enumerate(custom_products)
                        if item.get("catalog_id") == catalog_id
                    ),
                    None,
                )
                if existing_index is None:
                    custom_products.append(snapshot)
                else:
                    custom_products[existing_index] = snapshot

            results.append({"catalog_id": catalog_id, "is_custom": is_custom})

        write_json(PRODUCT_OVERRIDES_PATH, overrides)
        write_json(CUSTOM_PRODUCTS_PATH, custom_products)
        return {"message": "Products persisted", "count": len(results), "products": results}
    except HTTPException:
        raise
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/save")
async def save_draft(req: ProposalRequest):
    try:
        # 임시저장: 현재 작업 중인 선택 목록 전체를 초안으로 보관한다.
        saved_products = [p.model_dump() for p in req.products]
        write_json(DRAFT_PATH, saved_products)
        return {
            "message": "Draft saved successfully",
            "count": len(saved_products),
        }
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/load")
async def load_draft():
    try:
        draft = read_json(DRAFT_PATH, [])
        if not isinstance(draft, list):
            return []
        return draft
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/generate")
async def generate_proposal(req: ProposalRequest):
    try:
        prs = Presentation(TEMPLATE_PATH)  # 원본 템플릿 사용

        # 기존 템플릿 상품 vs 신규/커스텀 상품 분리
        existing_products = [p for p in req.products if not needs_slide_assignment(p)]
        new_products = [p for p in req.products if needs_slide_assignment(p)]

        selected_indices = [p.slide_index for p in existing_products]
        # 상품 슬라이드 범위: 6 ~ 마지막-1 (마지막은 THANK YOU)
        last_product_index = len(prs.slides) - 2
        unselected_indices = [
            i for i in range(6, last_product_index + 1)
            if i not in selected_indices
        ]

        # 새로운 상품에게 사용하지 않은 슬라이드 할당
        for p in new_products:
            if not unselected_indices:
                raise Exception("템플릿에 사용 가능한 여분 슬라이드가 부족합니다.")
            reused_idx = unselected_indices.pop(0)
            p.slide_index = reused_idx
            selected_indices.append(reused_idx)

        # 요청 순서 유지
        all_products = list(req.products)

        # 1. 선택된 슬라이드의 데이터 채우기
        for i, prod in enumerate(all_products):
            if prod.slide_index < 0 or prod.slide_index >= len(prs.slides):
                raise Exception(f"잘못된 슬라이드 인덱스: {prod.slide_index}")
            slide = prs.slides[prod.slide_index]
            is_custom_slide = needs_slide_assignment(prod) or bool(prod.is_new)

            # 상품명 / 상세설명 교체
            for idx, shape in enumerate(slide.shapes):
                if shape.has_text_frame and "KEY SELLING POINTS" in shape.text:
                    if idx > 0 and slide.shapes[idx - 1].has_text_frame:
                        if prod.name and prod.name.strip() != "":
                            replace_text_in_shape(
                                slide.shapes[idx - 1], prod.name, font_size=18, is_bold=True
                            )
                    if idx + 1 < len(slide.shapes) and slide.shapes[idx + 1].has_text_frame:
                        description = (prod.description or "").strip()
                        if description:
                            replace_text_in_shape(slide.shapes[idx + 1], description, font_size=10)
                        elif is_custom_slide:
                            replace_text_in_shape(slide.shapes[idx + 1], "• 상세 설명 참고", font_size=10)
                    break

            # 신규/커스텀 상품 이미지 교체
            if is_custom_slide and prod.image_base64:
                try:
                    img_shape = None
                    for s in slide.shapes:
                        if s.shape_type == 13:  # PICTURE
                            img_shape = s
                            break
                    if img_shape:
                        left, top, width, height = (
                            img_shape.left,
                            img_shape.top,
                            img_shape.width,
                            img_shape.height,
                        )
                        sp = img_shape.element
                        sp.getparent().remove(sp)

                        img_data = (
                            prod.image_base64.split(",")[1]
                            if "," in prod.image_base64
                            else prod.image_base64
                        )
                        img_bytes = base64.b64decode(img_data)
                        tmp_img_path = f"temp_img_{i}.png"
                        with open(tmp_img_path, "wb") as f:
                            f.write(img_bytes)

                        slide.shapes.add_picture(tmp_img_path, left, top, width, height)
                        if os.path.exists(tmp_img_path):
                            os.remove(tmp_img_path)
                except Exception as img_err:
                    print(f"이미지 교체 오류: {img_err}")

            # 공통 데이터: 미작성 항목은 공란, 작성 항목만 넣고 템플릿 잔여/중복 박스 제거
            apply_product_b2b_fields(slide, prod)

        # 2. 표지 + 선택 상품(요청 순서) + 마지막 장만 남기고 재배치
        #    설명 슬라이드(2~6)와 미선택 상품은 제외된다.
        thank_you_index = len(prs.slides) - 1
        ordered_indices = [0] + [p.slide_index for p in all_products] + [thank_you_index]
        # 중복 인덱스 제거(동일 슬라이드 이중 참조 방지)
        seen = set()
        unique_ordered = []
        for idx in ordered_indices:
            if idx in seen:
                continue
            seen.add(idx)
            unique_ordered.append(idx)
        rebuild_slides_in_order(prs, unique_ordered)

        output = io.BytesIO()
        prs.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": "attachment; filename=ESSER_Custom_Proposal.pptx"},
        )

    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


# Mount frontend dist folder
frontend_dist = os.path.join(os.path.dirname(__file__), "../frontend/dist")
if os.path.exists(frontend_dist):
    app.mount("/", StaticFiles(directory=frontend_dist, html=True), name="frontend")
