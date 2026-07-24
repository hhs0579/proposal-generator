import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs2 = Presentation("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/test_gen.pptx")
slide2 = prs2.slides[7]
print("All text in Slide 8 of generated PPT:")
for s in slide2.shapes:
    if s.has_text_frame and s.text.strip():
        print(repr(s.text.strip()))
