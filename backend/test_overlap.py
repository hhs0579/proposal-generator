import sys
from pptx import Presentation

# Generate a proposal with a known new value to see if the old value is still there
prs = Presentation("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/template.pptx")
slide = prs.slides[7] # Slide 8

# Force write "테스트 가격" to the supply price box
for idx, shape in enumerate(slide.shapes):
    if shape.has_text_frame and shape.text.strip() == "공급가":
        slide.shapes[idx+2].text = "테스트 가격"

prs.save("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/test_gen.pptx")

# Now load it and print all texts in the slide
prs2 = Presentation("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/test_gen.pptx")
slide2 = prs2.slides[7]
print("All text in Slide 8 of generated PPT:")
for s in slide2.shapes:
    if s.has_text_frame and s.text.strip():
        print(repr(s.text.strip()))
