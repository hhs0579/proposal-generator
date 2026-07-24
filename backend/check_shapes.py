from pptx import Presentation
prs = Presentation("template.pptx")
shapes = prs.slides[7].shapes
with open("shapes2.log", "w", encoding="utf-8") as f:
    for i, s in enumerate(shapes):
        t = s.text.replace("\n", " ") if s.has_text_frame else "NO_TEXT"
        f.write(f"{i} {s.name} {repr(t)}\n")
