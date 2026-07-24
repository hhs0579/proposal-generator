import sys
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/esser dev/proposal-generator/backend/test_overlap.pptx")
# Find the Options text box on Slide 8
slide = prs.slides[7]
for idx, shape in enumerate(slide.shapes):
    if shape.has_text_frame and "색상 | 블랙" in shape.text:
        print(shape.element.xml)
        break
