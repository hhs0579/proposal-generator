import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")
slide = prs.slides[6]

print(f"--- Slide 7 Shapes ---")
for idx, shape in enumerate(slide.shapes):
    print(f"Shape {idx}: type={shape.shape_type}")
    if shape.has_text_frame:
        print(f"  Text: {repr(shape.text)}")
    elif shape.has_table:
        print(f"  Table: {len(shape.table.rows)} rows")
        for r_idx, row in enumerate(shape.table.rows):
            row_data = [cell.text_frame.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"    Row {r_idx}: {row_data}")
