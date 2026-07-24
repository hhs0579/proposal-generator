import io
import sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")

for i, slide in enumerate(prs.slides):
    if i < 6 or i > 54:
        continue
    
    product_name = None
    for idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame and "KEY SELLING POINTS" in shape.text:
            if idx > 0 and slide.shapes[idx-1].has_text_frame:
                product_name = slide.shapes[idx-1].text.strip()
            break
            
    if not product_name:
        # Fallback if "KEY SELLING POINTS" not found
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip().replace('\n', ' ')
                if text and len(text) > 2 and len(text) < 50 and "상품 제안서" not in text and "PRODUCT LINE-UP" not in text:
                    product_name = text
                    break

    print(f"Slide {i+1}: {repr(product_name)}")
