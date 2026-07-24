import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation

prs = Presentation("C:/Users/Esser/Desktop/ESSER_B2B_상품제안서_수정본.pptx")
slide = prs.slides[7]
layout = slide.slide_layout
print(f"Slide Layout name: {layout.name}")
for idx, shape in enumerate(layout.shapes):
    if shape.has_text_frame:
        print(f"Layout Shape {idx}: {repr(shape.text.strip())}")
